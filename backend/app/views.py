import hashlib, re, secrets, difflib
from datetime import timedelta
from django.db import transaction, models
from django.utils import timezone
from django.core.validators import validate_email
from django.http import JsonResponse, HttpResponse, FileResponse
from django.core.mail import send_mail
from django.conf import settings
from django.views.decorators.csrf import ensure_csrf_cookie
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.contrib.auth.models import User
from rest_framework.decorators import api_view, permission_classes, parser_classes
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.response import Response
from rest_framework import status
from rest_framework.pagination import PageNumberPagination
from .models import Workspace, Membership, WorkspaceInvitation, Guideline, GuidelineVersion, GuidelineChunk, GuidelineReference, Rule, ApprovedExample, Asset, AssetVersion, Finding, Comment, Review, AuditEvent, ExtractedClaim, EvidenceSource, EvidenceReference, CorrectionSuggestion, BrandPaletteEntry, ReferenceLogo, AnalysisRun, GuidelineClarityInsight, FindingStatusEvent

def member(request, ws_id):
    return Membership.objects.filter(workspace_id=ws_id,user=request.user).select_related('workspace').first()
def wdata(w):
    return {'id':w.id,'name':w.name}
def maybe_paginate(request,data):
    if not any(key in request.query_params for key in ('page','page_size')):return Response(data)
    paginator=PageNumberPagination();paginator.page_size=25;paginator.page_size_query_param='page_size';paginator.max_page_size=100
    return paginator.get_paginated_response(paginator.paginate_queryset(data,request))
def boolean_value(value):
    if isinstance(value,bool):return value
    if isinstance(value,str) and value.lower() in {'true','1','yes','on'}:return True
    if isinstance(value,str) and value.lower() in {'false','0','no','off'}:return False
    return None
def ruledata(r):
    return {'id':r.id,'title':r.title,'text':r.text,'category':r.category,'severity':r.severity,'replacement':r.replacement,'active':r.active}
def sync_guideline_chunks(guideline):
    source=(guideline.source_text or '').strip()
    if not source:return []
    sections=re.split(r'(?=\[Page\s+\d+\])',source)
    chunks=[]
    for section in sections:
        section=section.strip()
        if not section:continue
        match=re.match(r'\[Page\s+(\d+)\]',section)
        location=f'Page {match.group(1)}' if match else ''
        body=section[match.end():].strip() if match else section
        paragraphs=[p.strip() for p in re.split(r'\n\s*\n',body) if p.strip()]
        current=''
        for paragraph in paragraphs:
            while len(paragraph)>1200:
                cut=paragraph.rfind(' ',0,1200)
                if cut<500:cut=1200
                part=paragraph[:cut].strip();paragraph=paragraph[cut:].strip()
                chunks.append((part,location))
            if current and len(current)+len(paragraph)+2>1200:
                chunks.append((current,location));current=''
            current=f'{current}\n\n{paragraph}'.strip()
        if current:chunks.append((current,location))
    return [GuidelineChunk.objects.create(guideline=guideline,version=guideline.version,index=i,text=text,location=location) for i,(text,location) in enumerate(chunks)]
def fdata(f):
    return {'id':f.id,'kind':f.kind,'severity':f.severity,'text':f.text,'explanation':f.explanation,'location':f.location,'status':f.status,'assigned_to':f.assigned_to.username if f.assigned_to_id else None,'rule':ruledata(f.rule) if f.rule else None,'claim':({'wording':f.claim.wording,'classification':f.claim.classification,'explanation':f.claim.explanation,'references':[{'source':x.source.name,'passage':x.passage,'location':x.location} for x in f.claim.references.select_related('source').all()],'guideline_references':[{'guideline':x.chunk.guideline.name,'version':x.chunk.version,'passage':x.passage,'location':x.location} for x in f.claim.guideline_references.select_related('chunk__guideline').all()]} if f.claim_id else None),'status_history':[{'from':x.previous,'to':x.current,'reason':x.reason,'actor':x.actor.username,'created_at':x.created_at} for x in f.status_events.select_related('actor').all()],'comments':[{'id':c.id,'body':c.body,'author':c.author.username,'parent_id':c.parent_id,'created_at':c.created_at,'replies':[{'id':reply.id,'body':reply.body,'author':reply.author.username,'created_at':reply.created_at} for reply in c.replies.select_related('author').all()]} for c in f.comments.filter(parent__isnull=True).select_related('author').prefetch_related('replies__author').all()],'suggestions':[{'minimal_edit':x.minimal_edit,'clearer_rewrite':x.clearer_rewrite,'rationale':x.rationale,'method':x.method,'scope':x.scope} for x in f.suggestions.all()]}
def vdata(v):
    return {'id':v.id,'number':v.number,'content':v.content,'extracted_text':v.extracted_text,'extraction_method':v.extraction_method,'extraction_status':v.extraction_status,'extraction_warnings':v.extraction_warnings,'extraction_metadata':v.extraction_metadata,'hash':v.content_hash,'created_at':v.created_at,'findings':[fdata(f) for f in v.findings.select_related('rule').prefetch_related('comments__author').all()]}
def adata(a):
    v=a.versions.order_by('-number').first()
    return {'id':a.id,'title':a.title,'campaign':a.campaign,'description':a.description,'status':a.status,'created_at':a.created_at,'latest_version':a.latest_version,'has_source_file':bool(a.source_file),'version':vdata(v) if v else None,'analysis_state':v.analysis_runs.order_by('-id').values_list('state',flat=True).first() if v else None,'warnings':v.analysis_runs.order_by('-id').values_list('warnings',flat=True).first() if v else []}
@api_view(['GET','POST'])
@permission_classes([AllowAny])
@ensure_csrf_cookie
def auth(request):
    if request.method=='GET': return Response({'authenticated':request.user.is_authenticated,'user':request.user.username if request.user.is_authenticated else None})
    action=request.data.get('action','login')
    if action=='register':
        username=request.data.get('username','').strip(); password=request.data.get('password','')
        if not username or len(password)<8: return Response({'error':'Username and password (8+ characters) are required.'},status=400)
        try:validate_password(password,User(username=username))
        except ValidationError as exc:return Response({'error':' '.join(exc.messages)},status=400)
        if User.objects.filter(username=username).exists(): return Response({'error':'Username already exists.'},status=400)
        token=request.data.get('invite_token','').strip();invitation=None
        if token:
            invitation=WorkspaceInvitation.objects.filter(token_hash=hashlib.sha256(token.encode()).hexdigest(),accepted_at__isnull=True,expires_at__gt=timezone.now()).first()
            if not invitation:return Response({'error':'Invitation link is invalid, expired or already used.'},status=400)
            if User.objects.filter(email__iexact=invitation.email).exists():return Response({'error':'This invited email already has an account. Ask an admin to add its username.'},status=409)
            if Membership.objects.filter(workspace=invitation.workspace,user__email__iexact=invitation.email).exists():return Response({'error':'An account with this invited email already belongs to the workspace.'},status=409)
        with transaction.atomic():
            user=User.objects.create_user(username=username,password=password,email=invitation.email if invitation else '')
            if invitation:
                Membership.objects.create(workspace=invitation.workspace,user=user,role=invitation.role)
                invitation.accepted_by=user;invitation.accepted_at=timezone.now();invitation.save(update_fields=['accepted_by','accepted_at'])
                AuditEvent.objects.create(workspace=invitation.workspace,actor=user,action='invitation_accepted',details={'email':invitation.email,'role':invitation.role})
        login(request,user)
        return Response({'authenticated':True,'user':username})
    identity=request.data.get('username','').strip()
    if '@' in identity:
        account=User.objects.filter(email__iexact=identity).order_by('id').first()
        identity=account.username if account else identity
    user=authenticate(request,username=identity,password=request.data.get('password'))
    if not user: return Response({'error':'Invalid username or password.'},status=400)
    login(request,user); return Response({'authenticated':True,'user':user.username})
@api_view(['POST'])
def auth_logout(request):
    logout(request); return Response({'ok':True})
@api_view(['GET','POST'])
def workspaces(request):
    if request.method=='GET': return Response([wdata(m.workspace) for m in Membership.objects.filter(user=request.user).select_related('workspace')])
    name=request.data.get('name','').strip()
    if not name:return Response({'error':'Workspace name is required.'},status=400)
    with transaction.atomic():
        w=Workspace.objects.create(name=name); Membership.objects.create(workspace=w,user=request.user,role='admin')
    return Response(wdata(w),status=201)

@api_view(['GET','PATCH'])
def workspace_settings(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if request.method=='GET':return Response({'retention_days':m.workspace.retention_days,'can_manage_retention':m.role=='admin'})
    if m.role!='admin':return Response({'error':'Workspace admin permission required.'},status=403)
    value=request.data.get('retention_days')
    if value in ('',None):m.workspace.retention_days=None
    else:
        try:value=int(value)
        except (ValueError,TypeError):return Response({'error':'Retention must be a positive number of days or null.'},status=400)
        if value<1 or value>3650:return Response({'error':'Retention must be between 1 and 3650 days.'},status=400)
        m.workspace.retention_days=value
    m.workspace.save(update_fields=['retention_days'])
    AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='retention_policy_updated',details={'retention_days':m.workspace.retention_days})
    return Response({'retention_days':m.workspace.retention_days,'can_manage_retention':True})
@api_view(['GET','POST'])
@parser_classes([MultiPartParser,FormParser,JSONParser])
def guidelines(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if request.method=='GET':return maybe_paginate(request,[{'id':g.id,'name':g.name,'version':g.version,'source_text':g.source_text,'rules':[ruledata(r) for r in g.rules.all()],'examples':[{'id':x.id,'text':x.text,'approved':x.approved} for x in g.examples.all()],'versions':[{'number':x.number,'created_at':x.created_at,'snapshot':x.snapshot} for x in g.versions.order_by('-number')]} for g in m.workspace.guidelines.all()])
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
    source=request.data.get('source_text','') or ''; upload=request.FILES.get('file'); warnings=[]
    if upload:
        if upload.size>20*1024*1024:return Response({'error':'Maximum guideline upload size is 20 MB.'},status=400)
        if upload.name.lower().endswith(('.png','.jpg','.jpeg')):return Response({'error':'Guidelines support TXT, PDF and DOCX files.'},status=400)
        try: source,warnings,_metadata=extract(upload)
        except ValueError as exc:return Response({'error':str(exc)},status=400)
        if not source.strip():return Response({'error':'No text could be extracted from this guideline document.'},status=400)
    name=(request.data.get('name','Brand rules') or '').strip()
    if not name:return Response({'error':'Guideline name is required.'},status=400)
    title=(request.data.get('rule_title','') or '').strip(); rule_text=(request.data.get('rule_text','') or '').strip()
    category=request.data.get('rule_category','general') or 'general'
    if bool(title)!=bool(rule_text):return Response({'error':'Both rule title and rule text are required.'},status=400)
    if title and category not in dict(Rule.CATEGORIES):return Response({'error':'Invalid rule category.'},status=400)
    with transaction.atomic():
        g=Guideline.objects.create(workspace=m.workspace,name=name,source_text=source,source_file=upload if upload else None)
        GuidelineVersion.objects.create(guideline=g,number=1,created_by=request.user,snapshot={'name':g.name,'source_text':g.source_text,'rules':[]})
        sync_guideline_chunks(g)
        created_rules=[]
        if title:
            rule=Rule.objects.create(guideline=g,title=title,text=rule_text,category=category,severity='medium',replacement=request.data.get('rule_replacement','') or '')
            g.version=2;g.save(update_fields=['version'])
            created_rules.append(ruledata(rule))
            GuidelineVersion.objects.create(guideline=g,number=2,created_by=request.user,snapshot={'name':g.name,'source_text':g.source_text,'rules':[ruledata(x) for x in g.rules.all()]})
            sync_guideline_chunks(g)
    return Response({'id':g.id,'name':g.name,'version':g.version,'source_text':g.source_text,'warnings':warnings,'rules':created_rules},status=201)
@api_view(['POST'])
def rules(request,ws_id,gid):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
    g=Guideline.objects.filter(id=gid,workspace=m.workspace).first()
    if not g:return Response({'error':'Guideline not found.'},status=404)
    title=request.data.get('title','').strip(); text=request.data.get('text','').strip(); category=request.data.get('category','general')
    if not title or not text:return Response({'error':'Rule title and text are required.'},status=400)
    if category not in dict(Rule.CATEGORIES):return Response({'error':'Invalid rule category.'},status=400)
    r=Rule.objects.create(guideline=g,title=title,text=text,category=category,severity=request.data.get('severity','medium'),replacement=request.data.get('replacement',''))
    number=max(g.version,g.versions.order_by('-number').values_list('number',flat=True).first() or 0)+1
    g.version=number;g.save(update_fields=['version'])
    GuidelineVersion.objects.create(guideline=g,number=number,created_by=request.user,snapshot={'name':g.name,'source_text':g.source_text,'rules':[ruledata(x) for x in g.rules.all()]})
    sync_guideline_chunks(g)
    return Response(ruledata(r),status=201)

@api_view(['PATCH'])
@parser_classes([MultiPartParser,FormParser,JSONParser])
def guideline_detail(request,ws_id,gid):
    m=member(request,ws_id);g=Guideline.objects.filter(id=gid,workspace_id=ws_id).first()
    if not m or not g:return Response({'error':'Guideline not found.'},status=404)
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
    name=request.data.get('name',g.name).strip();source=request.data.get('source_text',g.source_text);upload=request.FILES.get('file');warnings=[]
    if not name:return Response({'error':'Guideline name is required.'},status=400)
    if upload:
        if upload.size>20*1024*1024:return Response({'error':'Maximum guideline upload size is 20 MB.'},status=400)
        try:source,warnings,_metadata=extract(upload)
        except ValueError as exc:return Response({'error':str(exc)},status=400)
        if not source.strip():return Response({'error':'No text could be extracted; the existing guideline was not changed.'},status=400)
    before={'name':g.name,'source_text':g.source_text,'rules':[ruledata(x) for x in g.rules.all()]}
    with transaction.atomic():
        old_number=max(g.version,g.versions.order_by('-number').values_list('number',flat=True).first() or 0)
        if not g.versions.filter(number=old_number).exists():GuidelineVersion.objects.create(guideline=g,number=old_number,created_by=request.user,snapshot=before)
        g.name=name;g.source_text=source
        if upload:g.source_file=upload
        g.version=old_number+1;g.save()
        GuidelineVersion.objects.create(guideline=g,number=g.version,created_by=request.user,snapshot={'name':g.name,'source_text':g.source_text,'rules':[ruledata(x) for x in g.rules.all()]})
        sync_guideline_chunks(g)
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='guideline_updated',details={'guideline':g.id,'version':g.version})
    return Response({'id':g.id,'name':g.name,'version':g.version,'source_text':g.source_text,'warnings':warnings})

@api_view(['POST'])
def examples(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
    guideline=Guideline.objects.filter(id=request.data.get('guideline_id'),workspace=m.workspace).first()
    text=request.data.get('text','').strip();approved=request.data.get('approved')
    if not guideline:return Response({'error':'Guideline not found in this workspace.'},status=404)
    if not text or not isinstance(approved,bool):return Response({'error':'Example text and an approved boolean are required.'},status=400)
    example=ApprovedExample.objects.create(guideline=guideline,text=text,approved=approved)
    AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='guideline_example_added',details={'guideline':guideline.id,'example':example.id,'approved':approved})
    return Response({'id':example.id,'text':example.text,'approved':example.approved},status=201)
def _ocr_pil(image):
    import pytesseract
    from pytesseract import Output
    data=pytesseract.image_to_data(image,output_type=Output.DICT)
    regions=[];words=[]
    for i,raw in enumerate(data.get('text',[])):
        word=str(raw).strip()
        try:confidence=float(data['conf'][i])
        except (ValueError,TypeError,KeyError):confidence=-1
        if not word or confidence<0:continue
        x,y,w,h=(int(data[k][i]) for k in ('left','top','width','height'))
        if w<=0 or h<=0:continue
        words.append(word)
        regions.append({'text':word,'x':x,'y':y,'width':w,'height':h,'confidence':round(confidence,1)})
    confidence=round(sum(r['confidence'] for r in regions)/len(regions),1) if regions else None
    return ' '.join(words),regions,confidence

def extract(upload):
    name=upload.name.lower();ext=name.rsplit('.',1)[-1] if '.' in name else ''
    if ext in ('txt','md'):return upload.read().decode('utf-8',errors='replace'),[],{'method':'plain_text'}
    if ext=='pdf':
        try:
            from io import BytesIO
            from pypdf import PdfReader
            raw=upload.read();reader=PdfReader(BytesIO(raw));pages=[p.extract_text() or '' for p in reader.pages]
            metadata={'method':'pypdf','page_count':len(pages),'pages':[]};warnings=[]
            for index,text in enumerate(pages):metadata['pages'].append({'page':index+1,'method':'text' if text.strip() else 'empty'})
            empty=[index for index,text in enumerate(pages) if not text.strip()]
            if empty:
                try:
                    import fitz
                    import pytesseract
                    from PIL import Image
                    document=fitz.open(stream=raw,filetype='pdf')
                    for index in empty:
                        pdf_page=document[index]
                        pix=pdf_page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
                        rendered=Image.open(BytesIO(pix.tobytes('png'))).convert('RGB')
                        text,regions,confidence=_ocr_pil(rendered)
                        if text:
                            pages[index]=text;metadata['pages'][index]={'page':index+1,'method':'ocr','confidence':confidence,'regions':regions}
                        else:metadata['pages'][index]={'page':index+1,'method':'ocr','confidence':confidence,'regions':[]}
                    if any(metadata['pages'][index]['method']=='ocr' for index in empty):warnings.append('Scanned PDF pages were OCR processed; review OCR text and confidence before relying on it.')
                except Exception:
                    warnings.append('Some PDF pages contained no text. Scanned-page OCR is unavailable; install PyMuPDF and configure Tesseract.')
            return '\n'.join(f'[Page {i+1}]\n{text}' for i,text in enumerate(pages)),warnings,metadata
        except Exception as e:raise ValueError(f'PDF extraction failed: {e}')
    if ext=='docx':
        try:
            from docx import Document
            d=Document(upload);return '\n'.join(p.text for p in d.paragraphs),[],{'method':'python_docx','paragraph_count':len(d.paragraphs)}
        except Exception as e:raise ValueError(f'DOCX extraction failed: {e}')
    if ext in ('png','jpg','jpeg'):raise ValueError('Image uploads must use the image analysis pipeline.')
    raise ValueError('Supported file types: TXT, PDF, DOCX, PNG and JPEG.')

def analyze(v, ws):
    text=v.extracted_text
    for r in Rule.objects.filter(guideline__workspace=ws,active=True):
        if r.category=='prohibited' and r.text:
            for match in re.finditer(re.escape(r.text),text,re.I):
                f=Finding.objects.create(version=v,kind='terminology',severity=r.severity,text=match.group(),explanation=f'Configured rule “{r.title}”: {r.text}',rule=r)
                edit=text[:match.start()]+(r.replacement or '')+text[match.end():]
                clearer=edit
                if r.replacement:
                    clearer=re.sub(re.escape(r.text),r.replacement,text,flags=re.I)
                CorrectionSuggestion.objects.create(finding=f,minimal_edit=edit,clearer_rewrite=clearer,rationale=f'Removes the configured prohibited term {r.text!r}. Draft generated by deterministic rule substitution; review before applying.',scope='full_content')
        elif r.category=='preferred' and r.text and r.text.lower() in text.lower():
            f=Finding.objects.create(version=v,kind='preferred_term',severity=r.severity,text=r.text,explanation=f'Consider the configured preferred wording: {r.replacement or r.title}.',rule=r)
            edit=re.sub(re.escape(r.text),r.replacement or r.title,text,flags=re.I)
            CorrectionSuggestion.objects.create(finding=f,minimal_edit=edit,clearer_rewrite=edit,rationale=f'Uses the configured preferred wording from rule {r.title!r}. Deterministic substitution draft; review before applying.',scope='full_content')
        elif r.category=='claim' and re.search(r'\b(always|never|guaranteed|100%|#1|best)\b',text,re.I):
            Finding.objects.create(version=v,kind='claim_review',severity=r.severity,text='Absolute or promotional claim detected',explanation=f'Human review needed. Configured claim guidance: {r.text}',rule=r)
    from .ai import configured
    if not configured(): deterministic_claim_triage(v)

def deterministic_claim_triage(v):
    # No approved evidence means unverified, never contradicted.
    for sentence in re.split(r'(?<=[.!?])\s+',v.extracted_text):
        sentence=sentence.strip()
        if re.search(r'\d|\b(always|never|guaranteed|best|first|leading|proven)\b',sentence,re.I) and len(sentence)>15:
            claim=ExtractedClaim.objects.create(version=v,wording=sentence,classification='unverified',explanation='No approved evidence source was linked to this claim; human verification is required.')
            f=Finding.objects.create(version=v,kind='claim',severity='medium',text=sentence,explanation='Potential factual or promotional claim. No approved evidence source is configured, so this is unverified and requires human review.',claim=claim)
            CorrectionSuggestion.objects.create(finding=f,minimal_edit='[Remove or qualify this claim pending evidence.]',clearer_rewrite='[Add an approved source, or describe the benefit without an unsupported factual assertion.]',rationale='No relevant approved evidence was available. These are cautious templates, not AI-generated factual replacements.')

def run_configured_ai(v, ws):
    from .ai import analyze as provider_analyze, ProviderNotConfigured
    from django.db.models import F
    from .retrieval import retrieve_evidence, retrieve_evidence_with_embeddings, retrieve_guideline_chunks
    rules=list(Rule.objects.filter(guideline__workspace=ws,active=True))
    examples=list(ApprovedExample.objects.filter(guideline__workspace=ws).order_by('-created_at')[:20])
    evidence=retrieve_evidence(ws,v.extracted_text);guideline_chunks,retrieval_method=retrieve_guideline_chunks(ws,v.extracted_text);retrieval_warning=None
    embedding_model=__import__('os').getenv('LLM_EMBEDDING_MODEL')
    if embedding_model:
        from .ai import embed
        try:
            for source in ws.evidence_sources.filter(approved=True):
                if source.embedding_model!=embedding_model or not source.embedding:
                    source.embedding=embed(f'{source.name}\n{source.location}\n{source.passage}')
                    source.embedding_model=embedding_model
                    source.save(update_fields=['embedding','embedding_model'])
            current_chunks=GuidelineChunk.objects.filter(guideline__workspace=ws,version=F('guideline__version')).select_related('guideline')
            for chunk in current_chunks:
                if chunk.embedding_model!=embedding_model or not chunk.embedding:
                    chunk.embedding=embed(f'{chunk.guideline.name}\n{chunk.location}\n{chunk.text}')
                    chunk.embedding_model=embedding_model
                    chunk.save(update_fields=['embedding','embedding_model'])
            query_vector=embed(v.extracted_text)
            vector_sources,method=retrieve_evidence_with_embeddings(ws,v.extracted_text,query_vector,embedding_model)
            if method=='embeddings_cosine' and vector_sources:evidence=vector_sources
            else:retrieval_warning='No compatible approved evidence embeddings were available; lexical evidence retrieval was used.'
            guideline_chunks,guideline_method=retrieve_guideline_chunks(ws,v.extracted_text,query_vector,embedding_model)
            if guideline_method!='embeddings_cosine':retrieval_warning='Guideline-section embeddings were not available for this query; lexical guideline retrieval was used.'
        except Exception as exc:
            retrieval_warning=f'Embedding retrieval failed ({type(exc).__name__}); lexical evidence retrieval was used.'
            guideline_chunks,retrieval_method=retrieve_guideline_chunks(ws,v.extracted_text)
    try:
        result=provider_analyze(v.extracted_text,rules,evidence,guideline_chunks,examples)
    except ProviderNotConfigured as exc:
        return [str(exc)]
    except Exception as exc:
        deterministic_claim_triage(v)
        # Keep provider diagnostics useful without exposing response bodies, which
        # can contain request details. HTTPError status codes distinguish common
        # setup problems such as an invalid key, model, or rate limit.
        from urllib.error import HTTPError
        if isinstance(exc, HTTPError):
            status = exc.code
            if status in (401, 403):
                detail = f'Groq rejected the API key or its access (HTTP {status}). Check that LLM_API_KEY is an active Groq key.'
            elif status == 404:
                detail = 'Groq could not find the API route or model (HTTP 404). Check LLM_BASE_URL and LLM_MODEL.'
            elif status == 429:
                detail = 'Groq rate limit or usage quota reached (HTTP 429). Check Groq Console limits/billing and retry later.'
            elif status == 400:
                detail = 'Groq rejected the analysis request (HTTP 400). Check the model and request compatibility.'
            elif status >= 500:
                detail = f'Groq had a temporary server error (HTTP {status}). Retry the review later.'
            else:
                detail = f'Groq request failed (HTTP {status}). Check the provider configuration.'
            return [f'Configured AI analysis failed: {detail} Deterministic unverified-claim triage was retained.']
        return [f'Configured AI analysis failed ({type(exc).__name__}); deterministic unverified-claim triage was retained.']
    rule_map={r.id:r for r in rules}; evidence_map={e.id:e for e in evidence};chunk_map={chunk.id:chunk for chunk in guideline_chunks};warnings=[retrieval_warning] if retrieval_warning else []
    for row in result.get('claims',[])[:100]:
        if not isinstance(row,dict) or not isinstance(row.get('wording'),str) or not row['wording'].strip():
            warnings.append('Skipped malformed structured claim result.');continue
        classification=str(row.get('classification','Unverified')).lower()
        if classification not in {'supported','contradicted','unverified','uncertain'}: classification='uncertain'
        allowed_ids=[]
        evidence_ids=row.get('evidence_ids',[]);evidence_ids=evidence_ids if isinstance(evidence_ids,list) else []
        for source_id in evidence_ids:
            try: source=evidence_map.get(int(source_id))
            except (ValueError,TypeError): source=None
            if source: allowed_ids.append(source)
        if classification in {'supported','contradicted'} and not allowed_ids:
            classification='unverified'
            row['explanation']='No valid approved evidence reference was returned. '+str(row.get('explanation',''))
        allowed_chunks=[]
        chunk_ids=row.get('guideline_chunk_ids',[]);chunk_ids=chunk_ids if isinstance(chunk_ids,list) else []
        for chunk_id in chunk_ids:
            try:chunk=chunk_map.get(int(chunk_id))
            except (ValueError,TypeError):chunk=None
            if chunk:allowed_chunks.append(chunk)
        claim=ExtractedClaim.objects.create(version=v,wording=row['wording'][:10000],category=str(row.get('category','factual'))[:40],classification=classification,explanation=str(row.get('explanation',''))[:5000])
        for source in allowed_ids:
            EvidenceReference.objects.create(claim=claim,source=source,passage=source.passage,location=source.location)
        for chunk in allowed_chunks:GuidelineReference.objects.create(claim=claim,chunk=chunk,passage=chunk.text,location=chunk.location)
        allowed_rules=[]
        rule_ids=row.get('rule_ids',[]);rule_ids=rule_ids if isinstance(rule_ids,list) else []
        for rule_id in rule_ids:
            try:rule=rule_map.get(int(rule_id))
            except (ValueError,TypeError):rule=None
            if rule:allowed_rules.append(rule)
        f=Finding.objects.create(version=v,kind='claim',severity='medium',text=claim.wording,explanation=claim.explanation,claim=claim,rule=allowed_rules[0] if allowed_rules else None)
        minimal=str(row.get('minimal_edit','')).strip() or '[Qualify this claim or provide supporting evidence.]'
        rewrite=str(row.get('clearer_rewrite','')).strip() or minimal
        CorrectionSuggestion.objects.create(finding=f,minimal_edit=minimal,clearer_rewrite=rewrite,rationale='Suggested by the configured provider; verify source references and wording before applying.',method='configured LLM')
    for row in result.get('tone_findings',[])[:100]:
        if not isinstance(row,dict):continue
        try: rule=rule_map.get(int(row.get('rule_id')))
        except (ValueError,TypeError): rule=None
        if not rule or rule.category!='tone':continue
        Finding.objects.create(version=v,kind='tone',severity=rule.severity,text=str(row.get('affected_text',''))[:1000],explanation=str(row.get('explanation',''))[:5000],rule=rule)
    return warnings

def image_checks(upload, ws):
    warnings=[]
    try:
        from PIL import Image
        import pytesseract
        upload.seek(0);image=Image.open(upload).convert('RGB')
        text,regions,confidence=_ocr_pil(image)
        if not text:warnings.append('OCR returned no readable text.')
        elif confidence is not None and confidence<55:warnings.append('OCR confidence is low; review extracted text before analysis.')
        return text,warnings,{'method':'tesseract_ocr','image_size':{'width':image.width,'height':image.height},'ocr_confidence':confidence,'ocr_regions':regions}
    except Exception as exc:
        return '',[f'OCR unavailable or failed ({type(exc).__name__}); visual checks will still run.'],{'method':'tesseract_ocr','ocr_confidence':None,'ocr_regions':[]}

def run_visual_checks(version,workspace):
    warnings=[]
    asset=version.asset
    if not asset.source_file:return ['Visual analysis unavailable: original image was not retained.']
    try:
        import cv2, numpy as np
        from collections import Counter
        with asset.source_file.open('rb') as image_file:image_bytes=image_file.read()
        image=cv2.imdecode(np.frombuffer(image_bytes,dtype=np.uint8),cv2.IMREAD_COLOR)
        if image is None:return ['Visual analysis failed: image could not be decoded.']
        palette=list(BrandPaletteEntry.objects.filter(workspace=workspace))
        if palette:
            sample=cv2.resize(image,(128,128)).reshape(-1,3)
            colors=Counter(tuple(int(x) for x in c[::-1]) for c in sample).most_common(8)
            expected=[(p.name,p.hex_value,tuple(int(p.hex_value[i:i+2],16) for i in (1,3,5))) for p in palette]
            nearest=[]
            for color,count in colors[:5]:
                name,hx,rgb=min(expected,key=lambda p:sum((color[i]-p[2][i])**2 for i in range(3)))
                distance=round(sum((color[i]-rgb[i])**2 for i in range(3))**0.5,1)
                nearest.append({'detected':'#%02X%02X%02X'%color,'expected':name,'expected_hex':hx,'rgb_distance':distance,'pixel_share':round(count/len(sample),3)})
            Finding.objects.create(version=version,kind='visual_palette_advisory',severity='low',text='Dominant image colors compared with the configured palette.',explanation='Advisory whole-image color summary only. The system cannot isolate intended brand elements, so differences are not treated as violations.',location={'comparisons':nearest})
        else:warnings.append('No approved brand palette configured; color comparison unavailable.')
        logos=list(workspace.logos.all())
        if logos:
            gray=cv2.cvtColor(image,cv2.COLOR_BGR2GRAY);height,width=gray.shape
            for logo in logos:
                with logo.image.open('rb') as logo_file:logo_bytes=logo_file.read()
                reference=cv2.imdecode(np.frombuffer(logo_bytes,dtype=np.uint8),cv2.IMREAD_UNCHANGED)
                if reference is None:
                    warnings.append(f'Reference logo {logo.name!r} could not be decoded.');continue
                if len(reference.shape)==3 and reference.shape[2]==4:reference=reference[:,:,:3]
                ref=cv2.cvtColor(reference,cv2.COLOR_BGR2GRAY) if len(reference.shape)==3 else reference
                best=(-1,None)
                for scale in (0.2,0.3,0.4,0.5,0.65,0.8,1.0,1.2,1.4):
                    rw,rh=int(ref.shape[1]*scale),int(ref.shape[0]*scale)
                    if rw<12 or rh<12 or rw>width or rh>height:continue
                    tpl=cv2.resize(ref,(rw,rh)); result=cv2.matchTemplate(gray,tpl,cv2.TM_CCOEFF_NORMED);_,score,_,loc=cv2.minMaxLoc(result)
                    if score>best[0]:best=(score,(loc[0],loc[1],rw,rh))
                score,box=best
                if box is None:
                    warnings.append(f'Logo check for {logo.name!r} unsupported at this image size.');continue
                x,y,bw,bh=box;certain=score>=0.82
                Finding.objects.create(version=version,kind='visual_logo_match' if certain else 'visual_logo_uncertain',severity='low',text=f'Reference logo: {logo.name}',explanation=f'Basic grayscale template match confidence {score:.2f}. This check is scale-limited and does not reliably detect rotation, distortion, partial occlusion, or all logo placements; human review is required.',location={'x':x,'y':y,'width':bw,'height':bh,'confidence':round(float(score),3),'reference_logo_id':logo.id})
        else:warnings.append('No reference logo configured; logo matching unavailable.')
        regions=(version.extraction_metadata or {}).get('ocr_regions',[])
        contrast_checked=0
        def luminance(rgb):
            channels=[]
            for value in rgb:
                value=float(value)/255
                channels.append(value/12.92 if value<=0.04045 else ((value+0.055)/1.055)**2.4)
            return 0.2126*channels[0]+0.7152*channels[1]+0.0722*channels[2]
        for region in regions:
            if contrast_checked>=20 or region.get('confidence',0)<80:continue
            x=max(0,int(region.get('x',0)));y=max(0,int(region.get('y',0)));w=int(region.get('width',0));h=int(region.get('height',0))
            if w<8 or h<8 or x+w>width or y+h>height:continue
            crop=cv2.cvtColor(image[y:y+h,x:x+w],cv2.COLOR_BGR2RGB).reshape(-1,3)
            gray=np.dot(crop,[0.2126,0.7152,0.0722]);low=float(np.percentile(gray,15));high=float(np.percentile(gray,85))
            if high-low<35:continue
            foreground=crop[gray<=low];background=crop[gray>=high]
            if len(foreground)<4 or len(background)<4:continue
            fg=np.median(foreground,axis=0);bg=np.median(background,axis=0)
            l1,l2=sorted((luminance(fg),luminance(bg)),reverse=True);ratio=(l1+0.05)/(l2+0.05);contrast_checked+=1
            if ratio<4.5:
                Finding.objects.create(version=version,kind='visual_contrast_advisory',severity='low',text=f'Possible low text contrast near “{region.get("text","")[:80]}”.',explanation=f'Estimated contrast ratio {ratio:.2f}:1 from OCR word bounds and local luminance percentiles. Background segmentation is approximate; this is an advisory check, not accessibility certification.',location={'x':x,'y':y,'width':w,'height':h,'contrast_ratio':round(ratio,2),'foreground_rgb':[round(float(c)) for c in fg],'background_rgb':[round(float(c)) for c in bg],'ocr_confidence':region.get('confidence')})
        if not regions:warnings.append('Contrast check unavailable: OCR did not provide text regions.')
        elif not contrast_checked:warnings.append('No OCR text regions met the confidence and image-quality thresholds for contrast estimation.')
    except Exception as exc:
        warnings.append(f'Visual analysis failed ({type(exc).__name__}); other checks were retained.')
    return warnings

def process_analysis_run(run_id):
    """Execute one durable DB-backed analysis job; safe for worker retries."""
    with transaction.atomic():
        run=AnalysisRun.objects.select_for_update().select_related('version__asset','version__asset__workspace').get(id=run_id)
        if run.state not in ['queued','failed']:
            return run
        version=run.version;asset=version.asset;workspace=asset.workspace
        if version.analysis_runs.exclude(id=run.id).filter(state='failed').exists():
            version.findings.all().delete();version.claims.all().delete()
        run.state='processing';run.started_at=timezone.now();run.finished_at=None;run.errors=[];run.attempts+=1;run.save(update_fields=['state','started_at','finished_at','errors','attempts'])
        asset.status='processing';asset.save(update_fields=['status'])
    warnings=list(version.extraction_warnings or []);errors=[]
    filename=asset.source_file.name.rsplit('/',1)[-1] if asset.source_file else ''
    ext=filename.lower().rsplit('.',1)[-1] if '.' in filename else ''
    try:
        # A corrected version is already the user's edited analysis text. The
        # asset still retains its original upload for provenance/visual checks;
        # re-extracting it here would silently discard the correction.
        if asset.source_file and version.extraction_method!='user_edit':
            with asset.source_file.open('rb') as upload:
                if ext in ('png','jpg','jpeg'):
                    text,extraction_warnings,extraction_metadata=image_checks(upload,workspace)
                    method='tesseract_ocr'
                else:
                    text,extraction_warnings,extraction_metadata=extract(upload)
                    method={'txt':'plain_text','md':'plain_text','pdf':'pypdf','docx':'python_docx'}.get(ext,'unsupported')
            version.content=text;version.extracted_text=text;version.extraction_method=method;version.extraction_metadata={**(version.extraction_metadata or {}),**extraction_metadata}
            version.content_hash=hashlib.sha256(text.encode()).hexdigest()
            warnings.extend(extraction_warnings)
            if ext=='pdf':version.extraction_metadata={**(version.extraction_metadata or {}),'page_markers':'[Page N] markers embedded in extracted_text'}
        elif version.content:
            version.extracted_text=version.content
        else:
            warnings.append('No readable text was extracted; analysis cannot establish compliance.')
        version.extraction_status='warnings' if warnings else 'completed'
        version.extraction_warnings=warnings
        version.save(update_fields=['content','extracted_text','content_hash','extraction_method','extraction_status','extraction_warnings','extraction_metadata'])
    except Exception as exc:
        version.extraction_status='failed';version.extraction_warnings=[*warnings,f'Content extraction failed ({type(exc).__name__}); analysis did not establish compliance.']
        version.save(update_fields=['extraction_status','extraction_warnings'])
        errors.append(f'Content extraction failed ({type(exc).__name__}).')
    if not errors:
        if version.extracted_text.strip():
            try:analyze(version,workspace)
            except Exception as exc:errors.append(f'Rule and claim analysis failed ({type(exc).__name__}).')
            try:warnings.extend(run_configured_ai(version,workspace))
            except Exception as exc:warnings.append(f'AI analysis failed ({type(exc).__name__}); other results were retained.')
        else:
            warnings.append('No readable text was extracted; no textual compliance conclusion was made.')
        if ext in ('png','jpg','jpeg'):
            try:warnings.extend(run_visual_checks(version,workspace))
            except Exception as exc:warnings.append(f'Visual checks failed ({type(exc).__name__}); other analysis results were retained.')
    final_state='failed' if errors else 'warnings' if warnings else 'completed'
    run.state=final_state;run.warnings=warnings;run.errors=errors;run.finished_at=timezone.now();run.save(update_fields=['state','warnings','errors','finished_at'])
    asset.status=final_state;asset.save(update_fields=['status'])
    return run

@api_view(['GET','POST'])
@parser_classes([MultiPartParser,FormParser,JSONParser])
def assets(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if request.method=='GET':return maybe_paginate(request,[adata(a) for a in m.workspace.assets.order_by('-created_at')])
    title=request.data.get('title','').strip(); content=request.data.get('content',''); upload=request.FILES.get('file')
    if not title:return Response({'error':'Title is required.'},status=400)
    if upload:
        if upload.size>20*1024*1024:return Response({'error':'Maximum upload size is 20 MB.'},status=400)
        ext=upload.name.lower().rsplit('.',1)[-1] if '.' in upload.name else ''
        if ext not in {'txt','md','pdf','docx','png','jpg','jpeg'}:return Response({'error':'Supported file types: TXT, PDF, DOCX, PNG and JPEG.'},status=400)
    if not content.strip() and not upload:return Response({'error':'Paste content or select a supported file.'},status=400)
    with transaction.atomic():
        a=Asset.objects.create(workspace=m.workspace,title=title,campaign=request.data.get('campaign',''),description=request.data.get('description',''),created_by=request.user,source_file=upload if upload else None)
        ext=upload.name.lower().rsplit('.',1)[-1] if upload and '.' in upload.name else 'pasted_text'
        method={'txt':'plain_text','md':'plain_text','pdf':'pypdf','docx':'python_docx','png':'tesseract_ocr','jpg':'tesseract_ocr','jpeg':'tesseract_ocr'}.get(ext,'pasted_text')
        initial_content='' if upload else content
        v=AssetVersion.objects.create(asset=a,number=1,content=initial_content,extracted_text=initial_content,extraction_method=method,extraction_status='queued',extraction_metadata={'source_filename':upload.name if upload else None},content_hash=hashlib.sha256((initial_content or (upload.name if upload else '')).encode()).hexdigest(),created_by=request.user)
        AnalysisRun.objects.create(version=v,state='queued')
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='asset_submitted',details={'asset_id':a.id,'version':1})
    return Response(adata(a),status=202)

@api_view(['DELETE'])
def asset_detail(request,ws_id,asset_id):
    m=member(request,ws_id);a=Asset.objects.filter(id=asset_id,workspace_id=ws_id).first()
    if not m or not a:return Response({'error':'Asset not found.'},status=404)
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required to delete uploaded content.'},status=403)
    filename=a.source_file.name if a.source_file else None;storage=a.source_file.storage if a.source_file else None
    details={'asset_id':a.id,'title':a.title,'versions':a.versions.count()}
    with transaction.atomic():
        a.delete()
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='asset_deleted',details=details)
        if filename:transaction.on_commit(lambda:storage.delete(filename))
    return Response(status=204)

@api_view(['GET'])
def source_file_download(request,ws_id,kind,object_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    models_by_kind={'asset':Asset,'guideline':Guideline,'evidence':EvidenceSource}
    model=models_by_kind.get(kind)
    if not model:return Response({'error':'Unsupported source file type.'},status=404)
    item=model.objects.filter(id=object_id,workspace=m.workspace).first()
    if not item:return Response({'error':'Source file not found.'},status=404)
    upload=item.source_file
    if not upload:return Response({'error':'No original source file is stored for this item.'},status=404)
    try:
        filename=upload.name.rsplit('/',1)[-1].replace('\\','_').replace('"','_')
        response=FileResponse(upload.open('rb'),as_attachment=True,filename=filename)
    except (OSError,ValueError):return Response({'error':'The original source file is unavailable.'},status=410)
    response['Cache-Control']='private, no-store'
    AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='source_file_downloaded',details={'kind':kind,'object_id':item.id})
    return response
@api_view(['GET'])
def dashboard(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    assets_q=m.workspace.assets.all(); findings=Finding.objects.filter(version__asset__workspace=m.workspace)
    awaiting_ids=Review.objects.filter(version__asset__workspace=m.workspace,state='submitted',version__number=models.F('version__asset__latest_version')).values_list('version__asset_id',flat=True).distinct(); approved_ids=Review.objects.filter(version__asset__workspace=m.workspace,state='approved',version__number=models.F('version__asset__latest_version')).values_list('version__asset_id',flat=True).distinct()
    return Response({'total_assets':assets_q.count(),'awaiting_review':len(awaiting_ids)+assets_q.filter(status__in=['queued','processing']).exclude(id__in=awaiting_ids).count(),'open_findings':findings.filter(status='open').count(),'approved_assets':len(approved_ids),'recent':[adata(a) for a in assets_q.order_by('-created_at')[:5]]})
@api_view(['GET'])
def history(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    return maybe_paginate(request,[{'id':r.id,'asset':r.version.asset.title,'version':r.version.number,'state':r.state,'reviewer':r.reviewer.username,'reason':r.reason,'created_at':r.created_at} for r in Review.objects.filter(version__asset__workspace=m.workspace).select_related('version__asset','reviewer').order_by('-created_at')])
@api_view(['GET'])
def audit_history(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    rows=[{'id':e.id,'action':e.action,'actor':e.actor.username if e.actor_id else None,'details':e.details,'created_at':e.created_at} for e in AuditEvent.objects.filter(workspace=m.workspace).select_related('actor').order_by('-created_at')]
    return maybe_paginate(request,rows)

@api_view(['POST'])
def finding_action(request,ws_id,fid):
    m=member(request,ws_id); f=Finding.objects.filter(id=fid,version__asset__workspace_id=ws_id).first()
    if not m or not f:return Response({'error':'Finding not found.'},status=404)
    action=request.data.get('action')
    if action=='comment':
        body=request.data.get('body','').strip()
        if not body:return Response({'error':'Comment cannot be empty.'},status=400)
        parent=None
        if request.data.get('parent_id'):
            parent=Comment.objects.filter(id=request.data['parent_id'],finding=f).first()
            if not parent:return Response({'error':'Reply target must be a comment on this finding.'},status=400)
        Comment.objects.create(finding=f,author=request.user,body=body,parent=parent)
    elif action=='status':
        if request.data.get('status') not in dict(Finding.STATES):return Response({'error':'Invalid status.'},status=400)
        if request.data.get('status')=='dismissed' and not request.data.get('reason','').strip():return Response({'error':'A reason is required when dismissing a finding.'},status=400)
        old=f.status; f.status=request.data['status']; f.save(update_fields=['status']); FindingStatusEvent.objects.create(finding=f,actor=request.user,previous=old,current=f.status,reason=request.data.get('reason','')); AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='finding_status_changed',details={'finding':f.id,'from':old,'to':f.status,'reason':request.data.get('reason','')})
    elif action=='assign':
        if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required to assign reviewers.'},status=403)
        user_id=request.data.get('user_id')
        assignee=None
        if user_id:
            membership=Membership.objects.filter(workspace=m.workspace,user_id=user_id).select_related('user').first()
            if not membership:return Response({'error':'Assignee must be a member of this workspace.'},status=400)
            if membership.role not in ['admin','manager','reviewer']:return Response({'error':'Findings can only be assigned to a reviewer, manager or admin.'},status=400)
            assignee=membership.user
        previous=f.assigned_to.username if f.assigned_to_id else None;f.assigned_to=assignee;f.save(update_fields=['assigned_to'])
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='finding_assigned',details={'finding':f.id,'from':previous,'to':assignee.username if assignee else None})
    else:return Response({'error':'Unknown action.'},status=400)
    return Response(fdata(f))
@api_view(['POST'])
def review_action(request,ws_id,asset_id):
    m=member(request,ws_id); a=Asset.objects.filter(id=asset_id,workspace_id=ws_id).first()
    if not m or not a:return Response({'error':'Asset not found.'},status=404)
    state=request.data.get('state')
    if state not in dict(Review.STATES):return Response({'error':'Invalid review decision.'},status=400)
    if state in ['approved','rejected'] and m.role not in ['admin','manager','reviewer']:return Response({'error':'Reviewer permission required.'},status=403)
    v=a.versions.get(number=a.latest_version)
    reason=request.data.get('reason','').strip()
    if state in ['changes_requested','rejected'] and not reason:return Response({'error':'A reason is required for this review decision.'},status=400)
    if state=='approved' and v.findings.exclude(status__in=['resolved','dismissed']).exists() and not reason:
        return Response({'error':'Resolve open findings or record an approval override reason.'},status=400)
    with transaction.atomic():
        r=Review.objects.create(version=v,state=state,reviewer=request.user,reason=reason)
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='review_'+state,details={'asset':a.id,'version':v.number})
    return Response({'id':r.id,'state':r.state,'version':v.number},status=201)

@api_view(['POST'])
def retry_analysis(request,ws_id,asset_id):
    m=member(request,ws_id);a=Asset.objects.filter(id=asset_id,workspace_id=ws_id).first()
    if not m or not a:return Response({'error':'Asset not found.'},status=404)
    version=a.versions.filter(number=a.latest_version).first()
    if not version:return Response({'error':'No content version is available to retry.'},status=400)
    latest=version.analysis_runs.order_by('-id').first()
    if not latest or latest.state!='failed':return Response({'error':'Only failed analysis runs can be retried.'},status=409)
    with transaction.atomic():
        run=AnalysisRun.objects.create(version=version,state='queued')
        a.status='queued';a.save(update_fields=['status'])
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='analysis_retry_queued',details={'asset':a.id,'version':version.number,'run':run.id})
    return Response({'asset_id':a.id,'version_id':version.id,'analysis_run_id':run.id,'state':'queued'},status=202)
@api_view(['POST'])
def correction(request,ws_id,asset_id):
    m=member(request,ws_id); a=Asset.objects.filter(id=asset_id,workspace_id=ws_id).first()
    if not m or not a:return Response({'error':'Asset not found.'},status=404)
    content=request.data.get('content','')
    if not content.strip():return Response({'error':'Corrected content is required.'},status=400)
    with transaction.atomic():
        previous=a.versions.filter(number=a.latest_version).first()
        inherited_metadata=previous.extraction_metadata if previous else {}
        num=a.latest_version+1; v=AssetVersion.objects.create(asset=a,number=num,content=content,extracted_text=content,extraction_method='user_edit',extraction_status='completed',extraction_metadata=inherited_metadata,content_hash=hashlib.sha256(content.encode()).hexdigest(),created_by=request.user)
        a.latest_version=num; a.status='queued'; a.save(update_fields=['latest_version','status']); AnalysisRun.objects.create(version=v,state='queued')
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='version_created',details={'asset':a.id,'version':num})
    return Response(vdata(v),status=202)
@api_view(['GET'])
def receipt(request,ws_id,asset_id):
    m=member(request,ws_id); a=Asset.objects.filter(id=asset_id,workspace_id=ws_id).first()
    if not m or not a:return Response({'error':'Asset not found.'},status=404)
    requested=request.query_params.get('version')
    if requested:
        try:number=int(requested)
        except (TypeError,ValueError):return Response({'error':'Version must be a number.'},status=400)
        v=a.versions.filter(number=number).first()
    else:v=a.versions.filter(number=a.latest_version).first()
    if not v:return Response({'error':'Content version not found.'},status=404)
    reviews=list(v.reviews.select_related('reviewer').order_by('created_at'));review=reviews[-1] if reviews else None
    previous=a.versions.filter(number=v.number-1).first()
    change_diff=list(difflib.unified_diff((previous.content if previous else '').splitlines(),v.content.splitlines(),fromfile=f'version-{v.number-1}' if previous else 'empty',tofile=f'version-{v.number}',lineterm='')) if previous else []
    return Response({'asset':a.title,'asset_id':a.id,'workspace':m.workspace.name,'version_id':v.id,'version':v.number,'content_hash':v.content_hash,'reviewed_at':review.created_at if review else None,'reviewer':review.reviewer.username if review else None,'decision':review.state if review else 'No final decision recorded','review_history':[{'state':r.state,'reviewer':r.reviewer.username,'reason':r.reason,'created_at':r.created_at} for r in reviews],'correction_from_version':previous.number if previous else None,'corrections_applied':change_diff,'analysis_state':v.analysis_runs.order_by('-id').values_list('state',flat=True).first(),'warnings':v.extraction_metadata.get('warnings',[]) if isinstance(v.extraction_metadata,dict) else [],'findings':[fdata(f) for f in v.findings.all()],'guidelines':[{'name':g.name,'version':g.version,'rules':[ruledata(r) for r in g.rules.filter(active=True)]} for g in m.workspace.guidelines.all()],'sources_checked':[{'id':src.id,'name':src.name,'type':src.source_type,'location':src.location,'passage':src.passage} for src in EvidenceSource.objects.filter(id__in=EvidenceReference.objects.filter(claim__version=v).values_list('source_id',flat=True)).distinct()],'guideline_sections_checked':[{'guideline':ref.chunk.guideline.name,'version':ref.chunk.version,'location':ref.location,'passage':ref.passage} for ref in GuidelineReference.objects.filter(claim__version=v).select_related('chunk__guideline').distinct()],'note':'Only listed workspace-approved evidence references were used. A completed analysis does not imply compliance.'})

@api_view(['PATCH','DELETE'])
def rule_detail(request,ws_id,rid):
    m=member(request,ws_id); r=Rule.objects.filter(id=rid,guideline__workspace_id=ws_id).first()
    if not m or not r:return Response({'error':'Rule not found.'},status=404)
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
    if request.method=='DELETE':
        g=r.guideline; deleted=ruledata(r); r.delete(); number=max(g.version,g.versions.order_by('-number').values_list('number',flat=True).first() or 0)+1; g.version=number; g.save(update_fields=['version'])
        GuidelineVersion.objects.create(guideline=g,number=number,created_by=request.user,snapshot={'name':g.name,'source_text':g.source_text,'deleted_rule':deleted,'rules':[ruledata(x) for x in g.rules.all()]})
        sync_guideline_chunks(g)
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='guideline_rule_deleted',details={'rule':deleted,'version':number})
        return Response(status=204)
    before=ruledata(r)
    for field in ['title','text','category','severity','replacement','active']:
        if field in request.data:setattr(r,field,request.data[field])
    if not r.title.strip() or not r.text.strip():return Response({'error':'Rule title and text are required.'},status=400)
    if r.category not in dict(Rule.CATEGORIES):return Response({'error':'Invalid rule category.'},status=400)
    r.save()
    g=r.guideline;number=max(g.version,g.versions.order_by('-number').values_list('number',flat=True).first() or 0)+1;g.version=number;g.save(update_fields=['version'])
    GuidelineVersion.objects.create(guideline=g,number=number,created_by=request.user,snapshot={'name':g.name,'source_text':g.source_text,'rules':[ruledata(x) for x in g.rules.all()]})
    sync_guideline_chunks(g)
    AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='guideline_rule_updated',details={'guideline':g.id,'rule':r.id,'previous':before,'current':ruledata(r),'version':number})
    return Response(ruledata(r))

@api_view(['GET','POST'])
@parser_classes([MultiPartParser,FormParser,JSONParser])
def evidence(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if request.method=='POST':
        if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
        name=request.data.get('name','').strip();passage=request.data.get('passage','').strip();upload=request.FILES.get('file');warnings=[]
        if upload:
            if upload.size>20*1024*1024:return Response({'error':'Maximum evidence document size is 20 MB.'},status=400)
            try:passage,warnings,_metadata=extract(upload)
            except ValueError as exc:return Response({'error':str(exc)},status=400)
        approved=boolean_value(request.data.get('approved',False))
        if not name or not passage or approved is None:return Response({'error':'A source name, extractable passage and explicit approved boolean are required.'},status=400)
        e=EvidenceSource.objects.create(workspace=m.workspace,name=name,source_type='document' if upload else request.data.get('source_type','internal'),passage=passage,location=request.data.get('location',''),source_file=upload if upload else None,extraction_warnings=warnings,approved=approved)
        return Response({'id':e.id,'name':e.name,'source_type':e.source_type,'passage':e.passage,'location':e.location,'approved':e.approved,'warnings':warnings},status=201)
    return maybe_paginate(request,[{'id':e.id,'name':e.name,'source_type':e.source_type,'passage':e.passage,'location':e.location,'approved':e.approved} for e in m.workspace.evidence_sources.all()])

@api_view(['PATCH','DELETE'])
def evidence_detail(request,ws_id,eid):
    m=member(request,ws_id);e=EvidenceSource.objects.filter(id=eid,workspace_id=ws_id).first()
    if not m or not e:return Response({'error':'Evidence source not found.'},status=404)
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
    if request.method=='DELETE':
        if EvidenceReference.objects.filter(source=e).exists():return Response({'error':'This source is referenced by a review. Revoke approval to exclude it from future analysis while preserving its audit history.'},status=409)
        filename=e.source_file.name if e.source_file else None;storage=e.source_file.storage if e.source_file else None
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='evidence_source_deleted',details={'source_id':e.id,'name':e.name});e.delete()
        if filename:transaction.on_commit(lambda:storage.delete(filename))
        return Response(status=204)
    for field in ['name','source_type','passage','location']:
        if field in request.data:setattr(e,field,str(request.data[field]).strip())
    if 'approved' in request.data:
        value=request.data['approved']
        if not isinstance(value,bool):return Response({'error':'approved must be a boolean.'},status=400)
        e.approved=value
    if not e.name or not e.passage:return Response({'error':'Evidence name and passage are required.'},status=400)
    e.save();AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='evidence_source_updated',details={'source_id':e.id,'approved':e.approved})
    return Response({'id':e.id,'name':e.name,'source_type':e.source_type,'passage':e.passage,'location':e.location,'approved':e.approved})

@api_view(['GET','POST'])
def palette(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if request.method=='POST':
        if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
        color=request.data.get('hex_value','')
        if not re.fullmatch(r'#[0-9a-fA-F]{6}',color):return Response({'error':'Color must be a six-digit hex value.'},status=400)
        p=BrandPaletteEntry.objects.create(workspace=m.workspace,name=request.data.get('name','Brand color'),hex_value=color.upper())
        return Response({'id':p.id,'name':p.name,'hex_value':p.hex_value},status=201)
    return Response([{'id':p.id,'name':p.name,'hex_value':p.hex_value} for p in m.workspace.palette.all()])

@api_view(['GET'])
def claims(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    qs=ExtractedClaim.objects.filter(version__asset__workspace=m.workspace).select_related('version__asset').order_by('-created_at')
    if request.GET.get('classification'):qs=qs.filter(classification=request.GET['classification'])
    if request.GET.get('category'):qs=qs.filter(category__iexact=request.GET['category'])
    if request.GET.get('status'):qs=qs.filter(findings__status=request.GET['status'])
    if request.GET.get('severity'):qs=qs.filter(findings__severity=request.GET['severity'])
    if request.GET.get('q'):qs=qs.filter(wording__icontains=request.GET['q'].strip())
    return maybe_paginate(request,[{'id':c.id,'asset':c.version.asset.title,'version':c.version.number,'wording':c.wording,'location':c.location,'category':c.category,'classification':c.classification,'explanation':c.explanation,'references':[{'source':r.source.name,'passage':r.passage,'location':r.location} for r in c.references.select_related('source').all()],'guideline_references':[{'guideline':r.chunk.guideline.name,'version':r.chunk.version,'passage':r.passage,'location':r.location} for r in c.guideline_references.select_related('chunk__guideline').all()]} for c in qs])

@api_view(['GET','POST'])
def clarity(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if request.method=='GET':
        rows=[]
        for rule in Rule.objects.filter(guideline__workspace=m.workspace):
            fs=Finding.objects.filter(rule=rule)
            dismissed=fs.filter(status='dismissed').count(); changes=fs.filter(status='changes_requested').count(); discussions=Comment.objects.filter(finding__in=fs).count()
            n=dismissed+changes+discussions
            latest=rule.insights.order_by('-created_at').first()
            if n>=3 and not (latest and latest.state in ['accepted','edited','rejected'] and n<=latest.signal_count):rows.append({'rule_id':rule.id,'title':rule.title,'text':rule.text,'signals':{'dismissed':dismissed,'changes_requested':changes,'comments':discussions},'signal_count':n,'proposed_text':f'{rule.text} (Clarify the intended scope and include an approved example.)'})
        return Response(rows)
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
    rid=request.data.get('rule_id'); choice=request.data.get('decision')
    rule=Rule.objects.filter(id=rid,guideline__workspace=m.workspace).first()
    if not rule:return Response({'error':'Rule not found.'},status=404)
    if choice not in ['accept','edit','reject']:return Response({'error':'Decision must be accept, edit or reject.'},status=400)
    signal_findings=Finding.objects.filter(rule=rule)
    signal_count=signal_findings.filter(status='dismissed').count()+signal_findings.filter(status='changes_requested').count()+Comment.objects.filter(finding__in=signal_findings).count()
    if signal_count<3:return Response({'error':'At least three recorded review signals are required.'},status=400)
    signal_summary={'dismissed':signal_findings.filter(status='dismissed').count(),'changes_requested':signal_findings.filter(status='changes_requested').count(),'comments':Comment.objects.filter(finding__in=signal_findings).count()}
    proposed=f'{rule.text} (Clarify the intended scope and include an approved example.)'
    if choice=='reject':
        GuidelineClarityInsight.objects.create(guideline=rule.guideline,rule=rule,signal_count=signal_count,signal_summary=signal_summary,proposed_text=proposed,state='rejected')
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='clarity_suggestion_rejected',details={'rule':rule.id,'signal_count':signal_count})
        return Response({'decision':'rejected'})
    new_text=request.data.get('text') if choice=='edit' else proposed
    if not isinstance(new_text,str) or not new_text.strip():return Response({'error':'Proposed rule text cannot be empty.'},status=400)
    g=rule.guideline
    with transaction.atomic():
        old_no=g.versions.order_by('-number').values_list('number',flat=True).first() or g.version
        GuidelineVersion.objects.create(guideline=g,number=old_no+1,created_by=request.user,snapshot={'name':g.name,'rule_id':rule.id,'previous_text':rule.text,'updated_text':new_text})
        g.version=old_no+1;g.save(update_fields=['version']);rule.text=new_text;rule.save(update_fields=['text']);sync_guideline_chunks(g)
        GuidelineClarityInsight.objects.create(guideline=g,rule=rule,signal_count=signal_count,signal_summary=signal_summary,proposed_text=proposed,state='accepted' if choice=='accept' else 'edited')
        AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='guideline_rule_updated',details={'guideline':g.id,'rule':rule.id,'version':g.version,'source':'clarity_insight'})
    return Response({'decision':choice,'guideline_version':g.version,'rule':ruledata(rule)})

@api_view(['GET'])
def versions(request,ws_id,asset_id):
    m=member(request,ws_id);a=Asset.objects.filter(id=asset_id,workspace_id=ws_id).first()
    if not m or not a:return Response({'error':'Asset not found.'},status=404)
    rows=list(a.versions.order_by('number'))
    data=[{'id':v.id,'number':v.number,'content':v.content,'hash':v.content_hash,'created_at':v.created_at,'author':v.created_by.username,'approved':v.reviews.filter(state='approved').exists()} for v in rows]
    left=request.GET.get('from');right=request.GET.get('to')
    if left and right:
        from difflib import unified_diff
        old=a.versions.filter(number=left).first();new=a.versions.filter(number=right).first()
        if not old or not new:return Response({'error':'Version not found.'},status=404)
        data={'versions':data,'diff':list(unified_diff(old.content.splitlines(),new.content.splitlines(),fromfile=f'v{old.number}',tofile=f'v{new.number}',lineterm=''))}
    return Response(data)


@api_view(['GET'])
def integration_status(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    from .ai import configured
    try:
        import pytesseract
        pytesseract.get_tesseract_version(); ocr=True
    except Exception: ocr=False
    embedding_configured=bool(__import__('os').getenv('LLM_API_KEY') and __import__('os').getenv('LLM_EMBEDDING_MODEL'))
    return Response({'llm_configured':configured(),'llm_provider':'OpenAI-compatible' if configured() else None,'embeddings_configured':embedding_configured,'evidence_retrieval':'workspace-scoped embeddings cosine ranking with lexical fallback' if embedding_configured else 'workspace-scoped lexical ranking of approved passages','ocr_available':ocr,'visual_color_summary':True,'logo_matching':True,'contrast_analysis':True,'contrast_method':'advisory estimate from high-confidence OCR text boxes'})


@api_view(['POST'])
def adjudicate_claim(request,ws_id,claim_id):
    m=member(request,ws_id); claim=ExtractedClaim.objects.filter(id=claim_id,version__asset__workspace_id=ws_id).first()
    if not m or not claim:return Response({'error':'Claim not found.'},status=404)
    if m.role not in ['admin','manager','reviewer']:return Response({'error':'Reviewer permission required.'},status=403)
    decision=request.data.get('classification'); explanation=request.data.get('explanation','').strip()
    if decision not in dict(ExtractedClaim.CLASSES) or not explanation:return Response({'error':'A valid classification and explanation are required.'},status=400)
    source_id=request.data.get('source_id');source=None
    if source_id:
        source=EvidenceSource.objects.filter(id=source_id,workspace=m.workspace,approved=True).first()
        if not source:return Response({'error':'Only approved evidence from this workspace can be linked.'},status=400)
    if decision in ['supported','contradicted'] and not source:return Response({'error':'Supported or contradicted decisions require an approved evidence source.'},status=400)
    claim.classification=decision;claim.explanation=explanation;claim.save(update_fields=['classification','explanation'])
    if source:EvidenceReference.objects.create(claim=claim,source=source,passage=source.passage,location=source.location)
    AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='claim_adjudicated',details={'claim':claim.id,'classification':decision,'source_id':source.id if source else None})
    return Response({'id':claim.id,'classification':claim.classification,'explanation':claim.explanation})


@api_view(['GET','POST'])
@parser_classes([MultiPartParser,FormParser,JSONParser])
def logos(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if request.method=='GET':return Response([{'id':x.id,'name':x.name,'created_at':x.created_at} for x in m.workspace.logos.all()])
    if m.role not in ['admin','manager']:return Response({'error':'Brand manager permission required.'},status=403)
    image=request.FILES.get('image')
    if not image or image.size>5*1024*1024 or image.content_type not in ['image/png','image/jpeg']:
        return Response({'error':'Reference logo must be PNG or JPEG and no larger than 5 MB.'},status=400)
    logo=ReferenceLogo.objects.create(workspace=m.workspace,name=request.data.get('name','Reference logo'),image=image)
    return Response({'id':logo.id,'name':logo.name},status=201)


@api_view(['GET','POST'])
def members(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if request.method=='GET':return maybe_paginate(request,[{'id':x.id,'username':x.user.username,'email':x.user.email,'role':x.role} for x in m.workspace.members.select_related('user').order_by('user__username')])
    if m.role!='admin':return Response({'error':'Workspace admin permission required.'},status=403)
    username=request.data.get('username','').strip();role=request.data.get('role','creator')
    user=User.objects.filter(username=username).first()
    if not user:return Response({'error':'User must register before being added to this workspace.'},status=404)
    if role not in dict(Membership.ROLE):return Response({'error':'Invalid workspace role.'},status=400)
    item,created=Membership.objects.get_or_create(workspace=m.workspace,user=user,defaults={'role':role})
    if not created:return Response({'error':'User is already a workspace member.'},status=400)
    AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='member_added',details={'username':user.username,'role':role})
    return Response({'id':item.id,'username':user.username,'email':user.email,'role':item.role},status=201)

@api_view(['POST'])
def invitations(request,ws_id):
    m=member(request,ws_id)
    if not m:return Response({'error':'Workspace not found.'},status=404)
    if m.role!='admin':return Response({'error':'Workspace admin permission required.'},status=403)
    email=request.data.get('email','').strip().lower();role=request.data.get('role','creator')
    try:validate_email(email)
    except ValidationError:return Response({'error':'A valid email address is required.'},status=400)
    if role not in dict(Membership.ROLE):return Response({'error':'Invalid workspace role.'},status=400)
    if User.objects.filter(email__iexact=email).exists():return Response({'error':'This email already has an account. Add its username from the member list.'},status=409)
    if WorkspaceInvitation.objects.filter(workspace=m.workspace,email__iexact=email,accepted_at__isnull=True,expires_at__gt=timezone.now()).exists():return Response({'error':'An unexpired invitation already exists for this email.'},status=409)
    token=secrets.token_urlsafe(32)
    invitation=WorkspaceInvitation.objects.create(workspace=m.workspace,email=email,role=role,token_hash=hashlib.sha256(token.encode()).hexdigest(),created_by=request.user,expires_at=timezone.now()+timedelta(days=7))
    invite_path=f'/?invite={token}';email_sent=False;email_warning=None
    if getattr(settings,'EMAIL_HOST',''):
        try:
            send_mail(f'Invitation to {m.workspace.name} on BrandProof',f'{request.user.username} invited you to join {m.workspace.name} as {role}. Accept within seven days: {settings.PUBLIC_APP_URL}{invite_path}',settings.DEFAULT_FROM_EMAIL,[email],fail_silently=False)
            email_sent=True
        except Exception:
            email_warning='Invitation was created, but the email could not be sent. Share the invitation link manually.'
    AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='member_invited',details={'email':email,'role':role,'expires_at':invitation.expires_at.isoformat(),'email_sent':email_sent})
    return Response({'id':invitation.id,'email':email,'role':role,'expires_at':invitation.expires_at,'invite_path':invite_path,'email_sent':email_sent,'email_warning':email_warning},status=201)

@api_view(['PATCH','DELETE'])
def member_detail(request,ws_id,mid):
    m=member(request,ws_id);item=Membership.objects.filter(id=mid,workspace_id=ws_id).select_related('user').first()
    if not m or not item:return Response({'error':'Member not found.'},status=404)
    if m.role!='admin':return Response({'error':'Workspace admin permission required.'},status=403)
    if request.method=='DELETE':
        if item.role=='admin' and Membership.objects.filter(workspace_id=ws_id,role='admin').count()<2:return Response({'error':'A workspace must retain at least one admin.'},status=400)
        username=item.user.username;item.delete();AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='member_removed',details={'username':username});return Response(status=204)
    role=request.data.get('role')
    if role not in dict(Membership.ROLE):return Response({'error':'Invalid workspace role.'},status=400)
    if item.role=='admin' and role!='admin' and Membership.objects.filter(workspace_id=ws_id,role='admin').count()<2:return Response({'error':'A workspace must retain at least one admin.'},status=400)
    previous=item.role;item.role=role;item.save(update_fields=['role']);AuditEvent.objects.create(workspace=m.workspace,actor=request.user,action='member_role_changed',details={'username':item.user.username,'from':previous,'to':role})
    return Response({'id':item.id,'username':item.user.username,'email':item.user.email,'role':item.role})
